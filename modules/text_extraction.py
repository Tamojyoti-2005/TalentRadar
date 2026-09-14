
import io
import re

import pdfplumber
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]
    # also pull text out of tables (common in resume templates)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def extract_text(uploaded_file) -> str:
    """
    Accepts a Streamlit UploadedFile (or any object with .name and .read())
    and returns extracted plain text.
    """
    name = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    # reset pointer in case the caller re-reads this file later
    try:
        uploaded_file.seek(0)
    except Exception:
        pass

    if name.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        # assume plain text
        text = file_bytes.decode("utf-8", errors="ignore")

    return clean_text(text)


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
